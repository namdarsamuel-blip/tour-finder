#!/usr/bin/env python3
"""
Spotify Chart -> U.S. Tour Checker
===================================
1. Downloads every Spotify daily chart (Global + every country) from kworb.net,
   which mirrors Spotify's official charts as public pages.
2. Ranks every artist by how much they are charting (weighted toward the
   Global and U.S. charts).
3. Looks up each top artist on Ticketmaster (free Discovery API) and SeatGeek
   (free API) for upcoming U.S. shows: presale windows, on-sale dates, prices,
   links. SeatGeek also lists resale for most shows sold via AXS, DICE and
   venue box offices. Each artist also gets one-click AXS / Bandsintown /
   Songkick / Google search links for anything neither API covers.
4. Writes everything to one Excel file:
      Artists        - ranked artists + tour summary (dead / not-touring artists show 0 shows)
      US Shows       - one row per show, presales first, with ticket links
      Chart Entries  - every song on every chart (the full export)
      Charts         - which charts were pulled and their dates

Usage
-----
  pip install requests beautifulsoup4 pandas openpyxl
  python chart_tour_checker.py --tm-key YOUR_TICKETMASTER_KEY --sg-key YOUR_SEATGEEK_CLIENT_ID
  (either key alone also works)

  Useful options:
    --charts-only            just export the charts, skip the tour lookup
    --top 50                 only keep the top 50 songs of each chart (default: all 200)
    --countries us,global,ca only pull these charts (codes as in kworb URLs)
    --max-artists 300        how many top-ranked artists to check on Ticketmaster
    --include-featured       also credit featured artists (w/ ...)
    --weekly                 also pull weekly charts
    --out my_file.xlsx       output file name

  The key can also be set as an environment variable TICKETMASTER_API_KEY.
"""

import argparse
import datetime as dt
import os
import re
import sys
import time
import unicodedata
from collections import defaultdict

import pandas as pd
import requests
from bs4 import BeautifulSoup

KWORB = "https://kworb.net/spotify/"
TM_API = "https://app.ticketmaster.com/discovery/v2/"
HEADERS = {"User-Agent": "Mozilla/5.0 (chart-tour-checker)"}

# Chart weights for the popularity score
CHART_WEIGHT = {"global": 2.0, "us": 2.0}
DEFAULT_WEIGHT = 1.0

# Names that chart but are never real touring acts. Add your own here.
SKIP_ARTISTS = {
    "kpop demon hunters cast", "grand theft auto vi", "the hunger games",
    "the citizens of halloween", "various artists",
}


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def http_get(url, params=None, tries=5):
    for attempt in range(tries):
        try:
            r = requests.get(url, params=params, headers=HEADERS, timeout=30)
        except requests.RequestException as e:
            if attempt == tries - 1:
                raise
            time.sleep(2 + attempt * 2)
            continue
        if r.status_code == 429 or r.status_code >= 500:
            time.sleep(2 + attempt * 2)
            continue
        return r
    return r


def to_int(text):
    text = (text or "").strip().replace(",", "").replace("+", "")
    try:
        return int(text)
    except ValueError:
        return None


def norm(name):
    """Normalise an artist name for matching (accents, case, punctuation, &)."""
    s = unicodedata.normalize("NFKD", name or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.casefold().replace("&", "and")
    s = re.sub(r"^the\s+", "", s)
    return re.sub(r"[\W_]+", "", s)


def parse_iso(s):
    if not s:
        return None
    try:
        return dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return None


try:
    from zoneinfo import ZoneInfo
    _ET = ZoneInfo("America/New_York")
except Exception:  # pragma: no cover
    _ET = None


def fmt(d):
    if not d:
        return "?"
    return d.astimezone(_ET).strftime("%b %d %I:%M%p ET") if _ET else d.strftime("%b %d %I:%M%p UTC")


# --------------------------------------------------------------------------- #
# Part 1 - Spotify charts (via kworb.net)
# --------------------------------------------------------------------------- #
def list_charts(include_weekly=False, only=None):
    r = http_get(KWORB)
    r.encoding = "utf-8"
    soup = BeautifulSoup(r.text, "html.parser")
    charts, seen = [], set()
    for row in soup.find_all("tr"):
        cells = row.find_all("td")
        if not cells:
            continue
        country = cells[0].get_text(" ", strip=True)
        for a in row.find_all("a", href=True):
            m = re.search(r"country/([a-z]+)_(daily|weekly)\.html$", a["href"])
            if not m:
                continue
            code, period = m.groups()
            if period == "weekly" and not include_weekly:
                continue
            if only and code not in only:
                continue
            if (code, period) in seen:
                continue
            seen.add((code, period))
            charts.append({
                "code": code,
                "country": country,
                "period": period,
                "url": f"{KWORB}country/{code}_{period}.html",
            })
    return charts


def parse_chart_html(html, chart, top=None):
    soup = BeautifulSoup(html, "html.parser")
    m = re.search(r"\d{4}/\d{2}/\d{2}", soup.get_text(" ", strip=True)[:2000])
    chart_date = m.group(0).replace("/", "-") if m else ""

    table = soup.find("table")
    if table is None:
        return chart_date, []
    header_row = table.find("thead") or table.find("tr")
    headers = [th.get_text(" ", strip=True) for th in header_row.find_all(["th", "td"])]
    if "Artist and Title" not in headers:
        return chart_date, []
    idx = {h: i for i, h in enumerate(headers)}

    rows = []
    for tr in table.find_all("tr"):
        tds = tr.find_all("td")
        if len(tds) != len(headers):
            continue

        def col(name):
            i = idx.get(name)
            return tds[i].get_text(" ", strip=True) if i is not None else ""

        pos = to_int(col("Pos"))
        if pos is None:
            continue
        if top and pos > top:
            continue

        cell = tds[idx["Artist and Title"]]
        main, featured, title, track_id = [], [], "", ""
        after_track = False
        for a in cell.find_all("a", href=True):
            href = a["href"]
            ma = re.search(r"artist/([A-Za-z0-9]+)\.html", href)
            mt = re.search(r"track/([A-Za-z0-9]+)\.html", href)
            if mt:
                title, track_id, after_track = a.get_text(strip=True), mt.group(1), True
            elif ma:
                (featured if after_track else main).append((a.get_text(strip=True), ma.group(1)))
        if not main:  # plain-text fallback
            txt = cell.get_text(" ", strip=True)
            artist_txt, _, title = txt.partition(" - ")
            main = [(artist_txt.strip(), "")]

        change = col("P+")
        rows.append({
            "Chart": chart["country"],
            "Chart Code": chart["code"],
            "Period": chart["period"],
            "Chart Date": chart_date,
            "Pos": pos,
            "Change": change,
            "Artist": main[0][0],
            "Artist ID": main[0][1],
            "Featured": ", ".join(n for n, _ in featured),
            "_featured": featured,
            "Title": title,
            "Days on Chart": to_int(col("Days")),
            "Weeks on Chart": to_int(col("Wks") or col("Weeks")),
            "Peak": to_int(col("Pk")),
            "Streams": to_int(col("Streams")),
            "Streams Change": to_int(col("Streams+")),
            "Total Streams": to_int(col("Total")),
            "Spotify Track Link": f"https://open.spotify.com/track/{track_id}" if track_id else "",
        })
    return chart_date, rows


def listeners_page_url(page):
    return f"{KWORB}listeners.html" if page == 1 else f"{KWORB}listeners{page}.html"


def parse_listeners_html(html):
    """Rows of a kworb 'top artists by monthly listeners' page."""
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table")
    if table is None:
        return []
    heads = [th.get_text(" ", strip=True) for th in (table.find("thead") or table.find("tr")).find_all(["th", "td"])]
    idx = {h: i for i, h in enumerate(heads)}
    if "Listeners" not in idx:
        return []
    out = []
    for tr in table.find_all("tr"):
        tds = tr.find_all("td")
        if len(tds) != len(heads):
            continue
        a = tds[idx.get("Artist", 1)].find("a", href=True)
        m = re.search(r"artist/([A-Za-z0-9]+)", a["href"]) if a else None
        if not m:
            continue
        g = lambda h: to_int(tds[idx[h]].get_text(strip=True)) if h in idx else None
        out.append({"key": m.group(1), "name": a.get_text(strip=True), "rank": g("#"),
                    "listeners": g("Listeners"), "daily": g("Daily +/-"), "peak": g("PkListeners")})
    return out


def fetch_listeners(max_pages=20, delay=0.6):
    """Monthly listeners for every artist kworb tracks (about 25k artists above 500k listeners)."""
    out, seen = [], set()
    for page in range(1, max_pages + 1):
        r = http_get(listeners_page_url(page))
        if r.status_code != 200:
            break
        r.encoding = "utf-8"
        n_before = len(out)
        for row in parse_listeners_html(r.text):
            if row["key"] not in seen:
                seen.add(row["key"])
                out.append(row)
        print(f"  listeners page {page}: {len(out) - n_before} artists", flush=True)
        if len(out) == n_before:
            break
        time.sleep(delay)
    return out


def fetch_all_charts(charts, top=None, delay=0.6):
    entries, index = [], []
    for i, c in enumerate(charts, 1):
        print(f"  [{i}/{len(charts)}] {c['country']} {c['period']}", flush=True)
        r = http_get(c["url"])
        if r.status_code != 200:
            print(f"     ! skipped (HTTP {r.status_code})")
            continue
        r.encoding = "utf-8"
        date, rows = parse_chart_html(r.text, c, top)
        entries.extend(rows)
        index.append({"Chart": c["country"], "Code": c["code"], "Period": c["period"],
                      "Chart Date": date, "Songs": len(rows), "Source": c["url"]})
        time.sleep(delay)
    return entries, index


def rank_artists(entries, include_featured=False):
    stats = defaultdict(lambda: {
        "Artist": "", "Artist ID": "", "score": 0.0, "charts": set(), "songs": set(),
        "best_pos": 999, "best_chart": "", "us_best": None, "global_best": None,
        "top_song": "", "top_song_pos": 999, "new_songs": 0, "climbing": 0,
    })

    def credit(name, aid, e, weight):
        if not name or norm(name) in {norm(x) for x in SKIP_ARTISTS}:
            return
        key = aid or norm(name)
        s = stats[key]
        s["Artist"], s["Artist ID"] = name, aid
        w = CHART_WEIGHT.get(e["Chart Code"], DEFAULT_WEIGHT) * weight
        s["score"] += w * max(0, 201 - e["Pos"]) / 200
        s["charts"].add(e["Chart"])
        s["songs"].add(e["Title"])
        if e["Pos"] < s["best_pos"]:
            s["best_pos"], s["best_chart"] = e["Pos"], e["Chart"]
        if e["Chart Code"] == "us":
            s["us_best"] = min(filter(None, [s["us_best"], e["Pos"]]))
        if e["Chart Code"] == "global":
            s["global_best"] = min(filter(None, [s["global_best"], e["Pos"]]))
        if e["Chart Code"] in ("us", "global") and e["Pos"] < s["top_song_pos"]:
            s["top_song"], s["top_song_pos"] = e["Title"], e["Pos"]
        elif not s["top_song"]:
            s["top_song"] = e["Title"]
        if (e["Days on Chart"] or 9999) <= 60:
            s["new_songs"] += 1
        ch = e["Change"]
        if ch.startswith("+") or ch in ("NEW", "RE"):
            s["climbing"] += 1

    for e in entries:
        credit(e["Artist"], e["Artist ID"], e, 1.0)
        if include_featured:
            for n, aid in e["_featured"]:
                credit(n, aid, e, 0.5)

    rows = []
    for s in stats.values():
        rows.append({
            "Artist": s["Artist"],
            "Artist Key": s["Artist ID"] or norm(s["Artist"]),
            "Popularity Score": round(s["score"], 2),
            "# Charts": len(s["charts"]),
            "# Songs Charting": len(s["songs"]),
            "US Best Pos": s["us_best"],
            "Global Best Pos": s["global_best"],
            "Best Pos": s["best_pos"],
            "Best Chart": s["best_chart"],
            "Top Song": s["top_song"],
            "New Songs (<=60 days)": s["new_songs"],
            "Entries Climbing": s["climbing"],
            "Spotify Artist Link": f"https://open.spotify.com/artist/{s['Artist ID']}" if s["Artist ID"] else "",
        })
    df = pd.DataFrame(rows).sort_values("Popularity Score", ascending=False).reset_index(drop=True)
    df.insert(0, "Rank", range(1, len(df) + 1))
    return df


# --------------------------------------------------------------------------- #
# Part 2 - Tour lookup (Ticketmaster Discovery API)
# --------------------------------------------------------------------------- #
class Ticketmaster:
    def __init__(self, key):
        self.key = key
        self.calls = 0

    def get(self, path, **params):
        params["apikey"] = self.key
        for attempt in range(6):
            r = requests.get(TM_API + path, params=params, timeout=30)
            self.calls += 1
            time.sleep(0.22)  # stay under 5 requests/second
            if r.status_code == 429:
                time.sleep(2 + attempt * 2)
                continue
            if r.status_code in (401, 403):
                sys.exit("Ticketmaster rejected the API key (or daily quota of 5,000 calls is used up).")
            if r.status_code >= 500:
                time.sleep(2)
                continue
            if r.status_code != 200:
                return {}
            return r.json()
        return {}

    def find_artist(self, name):
        data = self.get("attractions.json", keyword=name, classificationName="music", size=20)
        atts = data.get("_embedded", {}).get("attractions", [])
        target = norm(name)
        exact = [a for a in atts if norm(a.get("name", "")) == target]
        return exact

    def us_events(self, attraction_id):
        out, page = [], 0
        while True:
            d = self.get("events.json", attractionId=attraction_id, countryCode="US",
                         size=200, page=page, sort="date,asc")
            out.extend(d.get("_embedded", {}).get("events", []))
            total_pages = d.get("page", {}).get("totalPages", 0)
            page += 1
            if page >= total_pages or page >= 4:
                break
        return out


class SeatGeek:
    """SeatGeek public API (free client_id from seatgeek.com/account/develop).
    SeatGeek lists primary + resale tickets, so it also surfaces many shows sold
    through AXS, DICE, Etix and independent venue box offices."""
    API = "https://api.seatgeek.com/2/"

    def __init__(self, client_id):
        self.cid = client_id
        self.calls = 0

    def get(self, path, **params):
        params["client_id"] = self.cid
        for attempt in range(5):
            try:
                r = requests.get(self.API + path, params=params, headers=HEADERS, timeout=30)
            except requests.RequestException:
                time.sleep(2)
                continue
            self.calls += 1
            time.sleep(0.25)
            if r.status_code == 429 or r.status_code >= 500:
                time.sleep(2 + attempt * 2)
                continue
            if r.status_code in (401, 403):
                sys.exit("SeatGeek rejected the client_id.")
            return r.json() if r.status_code == 200 else {}
        return {}

    def us_events(self, name):
        perf = self.get("performers", q=name, per_page=10).get("performers", [])
        target = norm(name)
        slugs = [p["slug"] for p in perf if norm(p.get("name", "")) == target]
        out = []
        for slug in slugs[:2]:
            page = 1
            while page <= 3:
                d = self.get("events", **{"performers.slug": slug, "venue.country": "US",
                                          "per_page": 100, "page": page, "sort": "datetime_local.asc"})
                evs = d.get("events", [])
                out.extend(evs)
                if len(evs) < 100:
                    break
                page += 1
        return out


def parse_sg_event(e, artist):
    v = e.get("venue") or {}
    st = e.get("stats") or {}
    when = e.get("datetime_local") or ""
    return {
        "Show ID": f"sg:{e.get('id', '')}",
        "Artist": artist,
        "Sale Status": "On sale (SeatGeek)",
        "Date": when[:10],
        "Time": when[11:16] if not e.get("time_tbd") else "",
        "Event": e.get("title", ""),
        "Venue": v.get("name", ""),
        "City": v.get("city", ""),
        "State": v.get("state", ""),
        "Public On-Sale": "",
        "Presales Open Now": "",
        "Upcoming Presales": "",
        "Price Min": st.get("lowest_price"),
        "Price Max": st.get("highest_price"),
        "Ticket Link": "",
        "SeatGeek Link": e.get("url", ""),
        "Source": "SeatGeek",
        "_sg_score": e.get("score"),
        "_sg_listings": st.get("visible_listing_count") or st.get("listing_count"),
        "_sg_low": st.get("lowest_price"),
        "_sg_median": st.get("median_price"),
        "_verified_fan": False,
        "_presale_count": 0,
    }


def search_links(name):
    from urllib.parse import quote_plus
    q = quote_plus(name)
    return {
        "AXS Search": f"https://www.axs.com/search?q={q}",
        "Bandsintown Search": f"https://www.bandsintown.com/search?query={q}",
        "Songkick Search": f"https://www.songkick.com/search?query={q}&type=artists",
        "Google Tour Search": f"https://www.google.com/search?q={quote_plus(name + ' tour 2026 2027 tickets USA')}",
    }


SALE_ORDER = {"PRESALE NOW": 0, "On sale now": 1, "On sale (SeatGeek)": 1, "Presale coming": 2,
              "Public on-sale coming": 3, "Rescheduled": 4, "Unknown": 5,
              "Not on sale (may be sold out)": 6, "Cancelled": 7}


def parse_event(e, artist, now):
    venue = (e.get("_embedded", {}).get("venues") or [{}])[0]
    dates = e.get("dates", {})
    start = dates.get("start", {})
    status = (dates.get("status", {}) or {}).get("code", "")
    sales = e.get("sales", {}) or {}
    pub_start = parse_iso((sales.get("public") or {}).get("startDateTime"))

    active, coming, windows = [], [], []
    for p in sales.get("presales", []) or []:
        s, en = parse_iso(p.get("startDateTime")), parse_iso(p.get("endDateTime"))
        if s and en and en > now:
            windows.append({"name": p.get("name", "Presale"), "start": s.isoformat(), "end": en.isoformat()})
        label = f"{p.get('name', 'Presale')} ({fmt(s)} - {fmt(en)})"
        if s and en and s <= now <= en:
            active.append(label)
        elif s and s > now:
            coming.append(label)

    # Ticketmaster says "offsale" whenever tickets aren't on sale right now, which
    # includes shows whose on-sale hasn't started yet, so sale dates win over it.
    if status in ("cancelled", "canceled"):
        sale = "Cancelled"
    elif active:
        sale = "PRESALE NOW"
    elif coming:
        sale = "Presale coming"
    elif pub_start and pub_start > now:
        sale = "Public on-sale coming"
    elif status == "offsale":
        sale = "Not on sale (may be sold out)"
    elif pub_start and pub_start <= now:
        sale = "On sale now"
    elif status == "rescheduled":
        sale = "Rescheduled"
    else:
        sale = "On sale now" if status == "onsale" else "Unknown"

    prices = e.get("priceRanges") or []
    pmin = min((p.get("min") for p in prices if p.get("min") is not None), default=None)
    pmax = max((p.get("max") for p in prices if p.get("max") is not None), default=None)

    return {
        "Show ID": f"tm:{e.get('id', '')}",
        "Artist": artist,
        "Sale Status": sale,
        "Date": start.get("localDate", ""),
        "Time": start.get("localTime", ""),
        "Event": e.get("name", ""),
        "Venue": venue.get("name", ""),
        "City": (venue.get("city") or {}).get("name", ""),
        "State": (venue.get("state") or {}).get("stateCode", ""),
        "Public On-Sale": fmt(pub_start) if pub_start else "",
        "Presales Open Now": "; ".join(active),
        "Upcoming Presales": "; ".join(coming),
        "Price Min": pmin,
        "Price Max": pmax,
        "Ticket Link": e.get("url", ""),
        "_presales": windows,
        "_onsale_at": pub_start.isoformat() if pub_start else None,
        "_status_code": status,
        "_verified_fan": any("verified" in (p.get("name") or "").lower() for p in sales.get("presales", []) or []),
        "_presale_count": len(sales.get("presales", []) or []),
        "_sg_score": None, "_sg_listings": None, "_sg_low": None, "_sg_median": None,
    }


def check_tours(artists_df, key, max_artists, sg_key=None):
    tm = Ticketmaster(key) if key else None
    sg = SeatGeek(sg_key) if sg_key else None
    now = dt.datetime.now(dt.timezone.utc)
    shows, summary = [], {}
    subset = artists_df.head(max_artists)
    for i, name in enumerate(subset["Artist"], 1):
        print(f"  [{i}/{len(subset)}] {name}", flush=True)
        events, tm_link = [], ""
        if tm:
            for att in tm.find_artist(name):
                upcoming = (att.get("upcomingEvents") or {}).get("_total", 1)
                tm_link = tm_link or att.get("url", "")
                if upcoming == 0:
                    continue
                events.extend(tm.us_events(att["id"]))
        seen, rows = set(), []
        for e in events:
            if e.get("id") in seen:
                continue
            seen.add(e.get("id"))
            row = parse_event(e, name, now)
            row["SeatGeek Link"] = ""
            row["Source"] = "Ticketmaster"
            rows.append(row)
        if sg:
            by_key = {(r["Date"], norm(r["City"])): r for r in rows}
            sg_seen = set()
            for e in sg.us_events(name):
                if e.get("id") in sg_seen:
                    continue
                sg_seen.add(e.get("id"))
                r = parse_sg_event(e, name)
                match = by_key.get((r["Date"], norm(r["City"])))
                if match:  # same show on both sites
                    match["SeatGeek Link"] = r["SeatGeek Link"]
                    match["Source"] = "Ticketmaster + SeatGeek"
                    for k in ("_sg_score", "_sg_listings", "_sg_low", "_sg_median"):
                        match[k] = r[k]
                    if match["Price Min"] is None:
                        match["Price Min"], match["Price Max"] = r["Price Min"], r["Price Max"]
                else:
                    rows.append(r)
                    by_key[(r["Date"], norm(r["City"]))] = r
        live = [r for r in rows if r["Sale Status"] not in ("Cancelled",)]
        shows.extend(rows)
        nxt = min(live, key=lambda r: r["Date"] or "9999") if live else None
        summary[name] = {
            "On Tour (US)": "YES" if live else "no",
            "US Shows": len(live),
            "Presale Now": sum(r["Sale Status"] == "PRESALE NOW" for r in live),
            "On Sale Now": sum(r["Sale Status"] == "On sale now" for r in live),
            "Presale/On-sale Coming": sum(r["Sale Status"] in ("Presale coming", "Public on-sale coming") for r in live),
            "On SeatGeek": sum(bool(r["SeatGeek Link"]) for r in live),
            "Next Show": f"{nxt['Date']} {nxt['City']}, {nxt['State']}" if nxt else "",
            "Next Show Date": nxt["Date"] if nxt else None,
            "States": ", ".join(sorted({r["State"] for r in live if r["State"]})),
            "Next Show Link": (nxt["Ticket Link"] or nxt["SeatGeek Link"]) if nxt else "",
            "Ticketmaster Artist Page": tm_link,
            **search_links(name),
        }
    if tm:
        print(f"  Ticketmaster calls used: {tm.calls} (free limit 5,000/day)")
    if sg:
        print(f"  SeatGeek calls used: {sg.calls}")
    return shows, summary


# --------------------------------------------------------------------------- #
# Excel output
# --------------------------------------------------------------------------- #
def write_excel(path, sheets):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2" if name == "Artists" else "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A5F")
            for ci, col in enumerate(df.columns, 1):
                letter = get_column_letter(ci)
                is_link = "Link" in col or "Page" in col or "Search" in col or col == "Source"
                q = df[col].astype(str).str.len().quantile(0.9) if len(df) else 10
                q = 10 if pd.isna(q) else int(q)
                width = 14 if is_link else min(60, max(10, q + 2, len(col) + 2))
                ws.column_dimensions[letter].width = width
                if is_link:
                    for r in range(2, ws.max_row + 1):
                        c = ws.cell(row=r, column=ci)
                        if isinstance(c.value, str) and c.value.startswith("http"):
                            c.hyperlink = c.value
                            c.value = "Open"
                            c.font = Font(color="0563C1", underline="single")
            if name == "US Shows" and "Sale Status" in df.columns:
                sc = list(df.columns).index("Sale Status") + 1
                green = PatternFill("solid", fgColor="C6EFCE")
                yellow = PatternFill("solid", fgColor="FFEB9C")
                for r in range(2, ws.max_row + 1):
                    v = ws.cell(row=r, column=sc).value
                    if v == "PRESALE NOW":
                        ws.cell(row=r, column=sc).fill = green
                    elif v in ("Presale coming", "Public on-sale coming"):
                        ws.cell(row=r, column=sc).fill = yellow
                    elif v == "On sale (SeatGeek)":
                        ws.cell(row=r, column=sc).fill = PatternFill("solid", fgColor="DDEBF7")


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description="Spotify charts -> U.S. tour checker")
    ap.add_argument("--tm-key", default=os.getenv("TICKETMASTER_API_KEY"))
    ap.add_argument("--sg-key", default=os.getenv("SEATGEEK_CLIENT_ID"), help="SeatGeek client_id")
    ap.add_argument("--charts-only", action="store_true")
    ap.add_argument("--top", type=int, default=None, help="songs per chart to keep (default all)")
    ap.add_argument("--countries", default="", help="comma list of chart codes, e.g. us,global,gb")
    ap.add_argument("--max-artists", type=int, default=250)
    ap.add_argument("--include-featured", action="store_true")
    ap.add_argument("--weekly", action="store_true")
    ap.add_argument("--out", default=f"chart_tour_report_{dt.date.today():%Y-%m-%d}.xlsx")
    a = ap.parse_args()

    only = {c.strip().lower() for c in a.countries.split(",") if c.strip()} or None

    print("Finding Spotify charts...")
    charts = list_charts(a.weekly, only)
    print(f"  {len(charts)} charts found. Downloading:")
    entries, index = fetch_all_charts(charts, a.top)
    if not entries:
        sys.exit("No chart data downloaded - check your internet connection.")

    artists = rank_artists(entries, a.include_featured)
    entries_df = pd.DataFrame(entries).drop(columns=["_featured"])
    sheets = {}

    if a.charts_only or not (a.tm_key or a.sg_key):
        if not a.charts_only:
            print("\nNo Ticketmaster or SeatGeek key given - exporting charts only.")
        sheets["Artists"] = artists
    else:
        print(f"\nChecking U.S. tour dates for the top {min(a.max_artists, len(artists))} artists...")
        shows, summary = check_tours(artists, a.tm_key, a.max_artists, a.sg_key)
        tour_cols = pd.DataFrame.from_dict(summary, orient="index")
        tour_cols.index.name = "Artist"
        merged = artists.merge(tour_cols, left_on="Artist", right_index=True, how="left")
        front = ["Rank", "Artist", "On Tour (US)", "US Shows", "Presale Now", "On Sale Now",
                 "Presale/On-sale Coming", "Next Show", "Next Show Link"]
        merged = merged[front + [c for c in merged.columns if c not in front]]
        sheets["Artists"] = merged
        shows_df = pd.DataFrame(shows)
        shows_df = shows_df[[col for col in shows_df.columns if not col.startswith("_")]]
        if not shows_df.empty:
            rank = dict(zip(artists["Artist"], artists["Rank"]))
            shows_df.insert(1, "Artist Rank", shows_df["Artist"].map(rank))
            shows_df["_o"] = shows_df["Sale Status"].map(SALE_ORDER).fillna(9)
            shows_df = shows_df.sort_values(["_o", "Artist Rank", "Date"]).drop(columns="_o")
        sheets["US Shows"] = shows_df

    sheets["Chart Entries"] = entries_df
    sheets["Charts"] = pd.DataFrame(index)
    write_excel(a.out, sheets)

    print(f"\nDone -> {os.path.abspath(a.out)}")
    if "US Shows" in sheets:
        on_tour = (sheets["Artists"]["On Tour (US)"] == "YES").sum()
        print(f"  {on_tour} of the checked artists have upcoming U.S. shows; "
              f"{len(sheets['US Shows'])} shows listed.")


if __name__ == "__main__":
    main()
